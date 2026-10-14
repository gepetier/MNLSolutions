"""Inspect Excel files without a Codex-specific runtime."""
import argparse
import json
from itertools import islice
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument("workbook", type=Path)
args = parser.parse_args()
if args.workbook.suffix.lower() == ".xls":
    import xlrd
    book = xlrd.open_workbook(args.workbook)
    try:
        for sheet in book.sheets():
            print(json.dumps({"sheet": sheet.name, "rows": sheet.nrows, "columns": sheet.ncols,
                              "preview": [sheet.row_values(i)[:12] for i in range(min(sheet.nrows, 30))]},
                             ensure_ascii=False, default=str))
    finally:
        book.release_resources()
else:
    from openpyxl import load_workbook
    book = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        for sheet in book:
            print(json.dumps({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                              "preview": list(islice(sheet.iter_rows(max_col=12, values_only=True), 30))},
                             ensure_ascii=False, default=str))
    finally:
        book.close()
