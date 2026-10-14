from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SENSITIVE_DATA_DIR = Path(os.environ.get("MNL_SENSITIVE_DATA_DIR", r"Z:\Adri\Dades MNLSolutions"))
PYDEPS = ROOT / "pydeps"
if PYDEPS.exists():
    sys.path.insert(0, str(PYDEPS))

QUOTE_PREFIXES = ("presp", "comanda", "pedido")
SKIP_PREFIXES = ("no valid", "no_val", "nou format")


def compact_spaces(value) -> str:
    text = str(value or "").strip()
    if "Ã" in text or "Â" in text:
        try:
            text = text.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    return re.sub(r"\s+", " ", text)


def document_ref_from_name(path: Path) -> str:
    match = re.search(r"MNL[-._ ]?(\d{2})[·.\-_\s]*(\d{4})[-._ ]?([A-Z])?", path.name, re.I)
    if not match:
        return ""
    suffix = f"-{match.group(3).upper()}" if match.group(3) else ""
    return f"MNL-{match.group(1)}·{match.group(2)}{suffix}"


def slug(value: str) -> str:
    value = compact_spaces(value).upper()
    value = value.replace("·", "-")
    value = re.sub(r"[^A-Z0-9]+", "-", value).strip("-")
    return value[:70] or "ARTICLE"


def parse_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = compact_spaces(value).replace(".", "").replace(",", ".") if "," in compact_spaces(value) else compact_spaces(value)
    try:
        return float(text)
    except ValueError:
        return None


def candidate_files(root: Path):
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}:
            continue
        name = path.name.lower().strip()
        if name.startswith("~$") or name.startswith(SKIP_PREFIXES):
            continue
        if name.startswith(QUOTE_PREFIXES):
            yield path


def read_sheet(path: Path):
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=False)
        sheet = book.sheet_by_index(0)

        def cell(row: int, col: int):
            if row >= sheet.nrows or col >= sheet.ncols:
                return ""
            value = sheet.cell_value(row, col)
            if isinstance(value, float) and value.is_integer():
                return int(value)
            return value

        return sheet.nrows, cell

    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    sheet = book.worksheets[0]

    def cell(row: int, col: int):
        return sheet.cell(row=row + 1, column=col + 1).value

    return sheet.max_row, cell


def extract_products(path: Path):
    try:
        row_count, cell = read_sheet(path)
    except Exception as error:
        return [], {"sourceFile": str(path), "error": str(error)}

    products = []
    ref = document_ref_from_name(path)
    for row in range(row_count):
        description = compact_spaces(cell(row, 0))
        if not re.match(r"^\d+\.\d+\.-", description):
            continue
        description = re.sub(r"^\d+\.\d+\.-\s*", "", description).strip()
        if not description:
            continue

        detail_parts = []
        next_row = row + 1
        while next_row < min(row_count, row + 4):
            detail = compact_spaces(cell(next_row, 0))
            if re.match(r"^\d+\.\d+\.-", detail):
                break
            if detail and not re.match(r"^(PORTES|RESUMEN|FORMA DE PAGO)", detail, re.I):
                detail_parts.append(detail)
            next_row += 1

        unit = compact_spaces(cell(row, 9)).lower() or "ud"
        unit_price = parse_number(cell(row, 10))
        amount = parse_number(cell(row, 11))
        quantity = parse_number(cell(row, 8)) or parse_number(cell(row, 7))
        if unit_price is None or unit_price <= 0:
            continue

        products.append(
            {
                "name": description,
                "description": "\n".join(detail_parts),
                "unit": unit,
                "price": round(unit_price, 4),
                "quantity": round(quantity, 4) if quantity is not None else None,
                "amount": round(amount, 2) if amount is not None else None,
                "quoteRef": ref,
                "sourceFile": str(path),
            }
        )
    return products, None


def merge_products(products):
    grouped = {}
    observations = defaultdict(list)
    for product in products:
        key = f"{slug(product['name'])}|{product['unit']}"
        observations[key].append(product)
        if key not in grouped:
            grouped[key] = {
                "code": slug(product["name"]),
                "name": product["name"],
                "description": product["description"],
                "unit": product["unit"],
                "price": product["price"],
                "selected": True,
                "sourceQuoteRefs": [],
                "sourceFiles": [],
                "observedPrices": [],
            }
        item = grouped[key]
        if not item.get("description") and product.get("description"):
            item["description"] = product["description"]
        if product.get("quoteRef") and product["quoteRef"] not in item["sourceQuoteRefs"]:
            item["sourceQuoteRefs"].append(product["quoteRef"])
        if product.get("sourceFile") and product["sourceFile"] not in item["sourceFiles"]:
            item["sourceFiles"].append(product["sourceFile"])
        price = product["price"]
        if price not in item["observedPrices"]:
            item["observedPrices"].append(price)

    for key, item in grouped.items():
        prices = [product["price"] for product in observations[key]]
        item["occurrences"] = len(observations[key])
        item["minObservedPrice"] = min(prices)
        item["maxObservedPrice"] = max(prices)
        item["observedPrices"] = sorted(item["observedPrices"])
    return sorted(grouped.values(), key=lambda item: item["name"].upper())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("--output", type=Path, default=SENSITIVE_DATA_DIR / "quote-products.json")
    args = parser.parse_args()

    all_products = []
    errors = []
    scanned = 0
    for path in candidate_files(args.root):
        scanned += 1
        products, error = extract_products(path)
        if error:
            errors.append(error)
        all_products.extend(products)

    merged = merge_products(all_products)
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceRoot": str(args.root),
        "scannedFiles": scanned,
        "rawLineCount": len(all_products),
        "productCount": len(merged),
        "errorCount": len(errors),
        "products": merged,
        "errors": errors,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: payload[k] for k in ["scannedFiles", "rawLineCount", "productCount", "errorCount"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
