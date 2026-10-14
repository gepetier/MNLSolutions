from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PYDEPS = ROOT / "pydeps"
if PYDEPS.exists():
    sys.path.insert(0, str(PYDEPS))


def read_xls(path: Path, max_rows: int, max_cols: int):
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=False)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(min(sheet.nrows, max_rows)):
            values = []
            has_any = False
            for c in range(min(sheet.ncols, max_cols)):
                value = sheet.cell_value(r, c)
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                text = str(value).strip()
                if text:
                    has_any = True
                values.append(text)
            if has_any:
                rows.append({"row": r + 1, "values": values})
        sheets.append({"name": sheet.name, "rows": rows})
    return sheets


def read_xlsx(path: Path, max_rows: int, max_cols: int):
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=max_rows, min_col=1, max_col=max_cols, values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                rows.append({"row": len(rows) + 1, "values": values})
        sheets.append({"name": sheet.title, "rows": rows})
    return sheets


def main():
    path = Path(sys.argv[1])
    max_rows = int(sys.argv[2]) if len(sys.argv) > 2 else 45
    max_cols = int(sys.argv[3]) if len(sys.argv) > 3 else 14
    if path.suffix.lower() == ".xls":
        sheets = read_xls(path, max_rows, max_cols)
    else:
        sheets = read_xlsx(path, max_rows, max_cols)
    print(json.dumps({"file": str(path), "sheets": sheets}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
