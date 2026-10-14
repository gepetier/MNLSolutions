import os
from pathlib import Path
from openpyxl import load_workbook

root = Path(os.environ["MNL_INVOICE_SOURCE_DIR"])
files = sorted(p for p in root.rglob("*.xlsx") if "factura" in p.name.lower() or "invoice" in p.name.lower())
for index in [26, 168, 221, 232, 405, 418, 642, 643]:
    book = load_workbook(files[index-1], read_only=True, data_only=True)
    ws = book.worksheets[0]
    print("\nFILE", index)
    for rn, row in enumerate(ws.iter_rows(min_row=16, max_row=min(ws.max_row, 65), values_only=True), 16):
        vals = [f"{chr(65+i)}:{str(v)[:65]}" for i, v in enumerate(row[:10]) if v is not None and str(v).strip()]
        if vals:
            print(rn, " | ".join(vals))
