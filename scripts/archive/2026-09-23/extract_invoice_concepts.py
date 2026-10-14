from __future__ import annotations

import os

import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(os.environ["MNL_INVOICE_SOURCE_DIR"])
OUT = Path(os.environ["MNL_SENSITIVE_DATA_DIR"]) / "invoice_concepts.json"


def text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def folded(value):
    value = unicodedata.normalize("NFKD", text(value))
    return "".join(c for c in value if not unicodedata.combining(c)).upper()


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def normalize_description(value):
    value = re.sub(r"^\s*\d+(?:[.,]\d+)*\s*[.:-]*\s*", "", text(value))
    value = re.sub(r"^(?:M2|M²|ML|UT|UD|UDS|U|PU)\s*[-.:]+\s*", "", value, flags=re.I)
    value = re.sub(r"\s+", " ", value).strip(" .-–:;")
    return value


def unit_code(raw, description):
    m = re.search(r"(?:^|\.)\s*(M2|M²|ML|UT|UDS|UD)\s*[.-]", folded(description))
    if m:
        return m.group(1).replace("M²", "M2")
    raw = folded(raw)
    return raw if raw in {"SM", "TM", "M2", "ML", "UT", "UDS", "UD"} else ""


def classify(desc, unit):
    s = folded(desc)
    if re.search(r"\b(PAGO A CUENTA|PAGAMENT A COMPTE|PAGAMENT A CONTA|ANTICIPO|BESTRETA)\b", s):
        return "Ajustos i bestretes", "Pagament a compte", "text explícit"
    if re.search(r"\b(ABONO|ABONAMENT|DESCUENTO|DESCOMPTE|BONIFIC)\w*", s):
        return "Ajustos i bestretes", "Abonaments i descomptes", "text explícit"
    if re.search(r"\b(PENALIZ|PENALIT|SUPLEMENT|RECARGO|CARGO (?:POR )?COLOR|INCREMENT(?:O)? (?:PER |POR )?COLOR|INCREMENT(?:O)? COST|INCREMENT(?:O)? ?< ?100)\w*", s):
        return "Ajustos i bestretes", "Recàrrecs", "text explícit"
    if re.search(r"\b(PORTES|TRANSPORT|PORTS|FLETE|FREIGHT|GRUA|GRUAJE)\w*", s):
        return "Transport i ports", "Transport", "text explícit"
    if re.search(r"\b(CALCUL|CALCULO|ENGINYER|INGENIER|MODELAT|MODELAD|VISAT|VISADO|PLANOLS|PLANOS|PROYECT|PROJECT|DESPECEJ|ESPECEJ|DELINEA|HONORAR|LLISTAT DE MATERIAL|LISTADO DE MATERIAL|ESTUDIO|ESTUDI|DESPLACAMENT|DIETA)\w*", s):
        return "Enginyeria i serveis", "Serveis tècnics", "text explícit"
    if re.search(r"\b(JUNTA ESTANCA|JUNTA DE ESTANQUEIDAD|SELLAD|SELLAT|SILICONA|CINTA BUT|CINTA SELL|MASTIC|K-STOP)\w*", s):
        return "Juntes i segellat", "Segellat", "text explícit"
    if re.search(r"\b(CARGOL|TORNILL|TIRAFON|ARANDEL|REMACHE|REBLON|FIJACION|FIXACIO|FIXACION|TAPAJUNTA)\w*", s):
        return "Cargoleria i accessoris", "Fixacions i accessoris", "text explícit"
    if re.search(r"\b(CANALON|CANALO|CANAL OBRA|CANAL EXTERIOR|CANALETA|BAJANTE|BAIXANT|TAPA CANAL)\w*", s):
        return "Canalons i drenatge", "Canalons", "text explícit"
    if re.search(r"\b(OMEGA|OMEGAS|SUPORT|SOPORTE|SUBESTRUCTURA|ESTRUCTURA SECUNDARIA|CORREA|CORRETJA|PERFIL SANITARIO|PERFIL SANITARI|PERFIL DE SUSPENSION|PERFIL DE SUSPENSIO)\w*", s):
        return "Perfils i subestructures", "Perfils i suports", "text explícit"
    if re.search(r"\b(REMAT|REMATE|CUMBRERA|CAVALLET|CORNISA|RACONER|RINCONER|TAPETA|BABERO|VORERA|BORDON|CHAPA LISA|XAPA LLISA|XAPA LISA|CHAPA PLEGADA|XAPA PLEGADA)\w*", s):
        if re.search(r"\b(CHAPA|XAPA)\b", s):
            return "Remats i xapes", "Xapa", "text explícit"
        return "Remats i xapes", "Remat", "text explícit"
    if re.search(r"\b(LLUCERNARI|LUCERNARIO|LUCERNARI|CLARABOYA)\w*", s):
        return "Llucernaris", "Llucernari", "text explícit"
    if re.search(r"\b(PUERTA|PORTA SECCIONAL|PORTA TALLAFOC|PORTA PEATONAL|PORTA RAPIDA)\w*", s):
        return "Portes i accessos", "Portes", "text explícit"
    # Catàleg vigent: HI-CT, TZ-CR, HI-XT = coberta; HI-ST* = façana.
    if re.search(r"\b(HI[- ]?(?:PIR|PIRM|QUADCORE)[- ]?(?:CT|XT)|HI[- ]?(?:CT|XT)|TZ[- ]?CR|CUBIERTA|COBERTA|PANEL TEJA|PANELL TEULA)\b", s):
        model = "HI-CT" if "CT" in s else "TZ-CR" if "TZ" in s else "HI-XT" if "XT" in s else "Altres de coberta"
        return "Panells de coberta", model, "catàleg/text"
    if re.search(r"\b(HI[- ]?(?:PIR|PIRM)[- ]?(?:F|ST[MLXS]?)|HI[- ]?ST[MLXS]?|FACHADA|FACANA|FAÇANA|FRIGORIFIC|CAMARA FRIGORIFICA|INDUSPANEL PF ?600)\b", s):
        model = "HI-ST" if "HI-" in s and "ST" in s else "HI-F" if "HI-PIR-F" in s else "PF600" if "PF" in s else "Altres de façana"
        return "Panells de façana", model, "catàleg/text"
    if re.search(r"\b(PANEL|PANELL|LANA DE ROCA|POLICARBONAT|POLICARBONATO|POLIESTER|POLYESTER)\b", s):
        return "Altres panells", "Panell no identificat", "text genèric"
    if re.match(r"^\W*(COLOR|ACABAT|ACABADO)\b", s):
        return "Acabats i colors", "Color o acabat", "text explícit"
    return "Per revisar", "Sense classificar", "sense correspondència clara"


def refine_group(desc, family, group):
    s = folded(desc)
    if family == "Remats i xapes":
        if group == "Xapa":
            if "PERFOR" in s:
                return "Xapa perforada"
            if "LLISA" in s or "LISA" in s:
                return "Xapa llisa"
            return "Altres xapes"
        for pattern, label in [
            (r"CUMBRERA|CAVALLET", "Remat de carener"),
            (r"LATERAL|COSTAT", "Remat lateral"),
            (r"FRONTAL|TAPA FRONTAL", "Remat frontal"),
            (r"PIE DE PANEL|PEU DE PANELL|PEU DE PANEL", "Remat de peu de panell"),
            (r"CORONACION|CORONACIO", "Remat de coronació"),
            (r"ESQUINA|RINCON|RACO", "Remat de cantonada"),
            (r"TROQUEL", "Remat troquelat"),
            (r"PERIMETR", "Remat perimetral"),
        ]:
            if re.search(pattern, s):
                return label
        return "Altres remats"
    if family == "Transport i ports":
        if re.search(r"CON GRUA|AMB GRUA|GRUA PARA|GRUA PER|GRUAJE", s):
            return "Transport amb grua"
        if re.search(r"SIN GRUA|SENSE GRUA", s):
            return "Transport sense grua"
        if re.search(r"COMBINAD", s):
            return "Transport combinat"
        return "Altres transports"
    if family == "Perfils i subestructures":
        if re.search(r"CORREA|CORRETJA", s):
            return "Corretges"
        if re.search(r"PIEZA DE SOPORTE|PIEZA DE SOIORTE|PIEZAS INFERIO|SUPORT.*PANELL|SOPORTE.*PANEL", s):
            return "Suports de panell"
        if re.search(r"OMEGA", s):
            if re.search(r"INTERMEDI", s):
                return "Omeges intermèdies"
            if re.search(r"VERTICAL|SEPARAD", s):
                return "Omeges separadores"
            if re.search(r"PRINCIPAL|ESQUINA", s):
                return "Omeges principals"
            return "Altres omegues"
        return "Altres perfils i estructures"
    if family == "Enginyeria i serveis":
        if re.search(r"VISAT|VISADO", s):
            return "Visats"
        if re.search(r"MODELAT|MODELAD|PLANOLS|PLANOS|DESPECEJ|ESPECEJ|DELINEA", s):
            return "Modelatge i plànols"
        if re.search(r"DESPLACAMENT|DIETA", s):
            return "Desplaçaments i dietes"
        if re.search(r"CALCUL|CALCULO|ESTUDI|ESTUDIO", s):
            return "Càlculs i estudis"
        return "Altres serveis tècnics"
    if family == "Ajustos i bestretes" and group == "Recàrrecs":
        if "COLOR" in s or "ACABAT" in s:
            return "Recàrrec de color"
        if "100" in s:
            return "Recàrrec de comanda petita"
    return group


def main():
    files = sorted(p for p in ROOT.rglob("*.xlsx") if ("factura" in p.name.lower() or "invoice" in p.name.lower()) and not p.name.startswith("~$"))
    lines = []
    issues = []
    header_patterns = Counter()
    invoices_without_lines = []
    for index, path in enumerate(files, 1):
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book.worksheets[0]
            rows = list(sheet.iter_rows(values_only=True))
            invoice_ref = next((m.group(0) for row in rows[:10] for v in row if (m := re.search(r"\b[A-Z]\d{2}[-/]\d+\b", text(v), re.I))), "")
            head = None
            for i, row in enumerate(rows[:40]):
                cells = [folded(v) for v in row]
                desc_candidates = [j for j, c in enumerate(cells) if "DESCRIP" in c]
                amount_candidates = [j for j, c in enumerate(cells) if c in {"IMPORTE", "IMPORT", "TOTAL"}]
                if desc_candidates and amount_candidates and max(amount_candidates) > min(desc_candidates):
                    head = (i, min(desc_candidates), max(amount_candidates))
                    break
            if head is None:
                issues.append({"invoice": index, "issue": "capçalera no trobada"})
                continue
            hrow, dcol, acol = head
            header_patterns[(hrow + 1, dcol + 1, acol + 1)] += 1
            count = 0
            for ri in range(hrow + 1, len(rows)):
                row = rows[ri]
                cells = [text(v) for v in row]
                row_text = " ".join(cells)
                # Summary begins after the commercial lines; keeps PORTES above it.
                if re.search(r"\b(SUBTOTAL|BASE IMPONIBLE|BASE IMPOSABLE)\b", folded(row_text)):
                    break
                amount = row[acol] if acol < len(row) else None
                if is_number(amount) and any("TOTAL" in folded(v) for v in row[dcol+1:acol]):
                    break
                if not is_number(amount):
                    continue
                description = text(row[dcol] if dcol < len(row) else None)
                raw_code = text(row[0] if dcol > 0 else "")
                if not description:
                    # Some newer invoices put PORTES in C and the description in D.
                    candidates = []
                    for candidate in range(dcol + 1, acol):
                        value = text(row[candidate] if candidate < len(row) else None)
                        if value and not is_number(row[candidate]) and not re.fullmatch(r"[\d., €%-]+", value):
                            candidates.append(value)
                    description = candidates[-1] if candidates else ""
                if not description:
                    # Page counters and unlabeled totals are not billed concepts.
                    continue
                if re.search(r"\b(IVA|TOTAL|SUBTOTAL|FORMA DE PAGO)\b", folded(description)):
                    continue
                normalized = normalize_description(description)
                if not normalized:
                    continue
                unit = unit_code(raw_code, description)
                family, group, basis = classify(normalized, unit)
                group = refine_group(normalized, family, group)
                lines.append({
                    "invoice": index,
                    "invoice_ref": invoice_ref,
                    "row": ri + 1,
                    "description": description,
                    "concept": normalized,
                    "unit_code": unit,
                    "family": family,
                    "group": group,
                    "basis": basis,
                    "amount": round(float(amount), 4),
                })
                count += 1
            if count == 0:
                advance = next(((ri, row) for ri, row in enumerate(rows) if any("PAGO A CUENTA" in folded(v) or "PAGAMENT A COMPTE" in folded(v) for v in row)), None)
                if advance:
                    ri, row = advance
                    description = next(text(v) for v in row if "PAGO A CUENTA" in folded(v) or "PAGAMENT A COMPTE" in folded(v))
                    lines.append({"invoice": index, "invoice_ref": invoice_ref, "row": ri + 1, "description": description, "concept": "Pagament a compte de pressupost", "unit_code": "", "family": "Ajustos i bestretes", "group": "Pagament a compte", "basis": "text explícit", "amount": 0})
                else:
                    invoices_without_lines.append(index)
            book.close()
        except Exception as error:
            issues.append({"invoice": index, "issue": type(error).__name__ + ": " + str(error)[:100]})
        if index % 100 == 0:
            print("processed", index, "of", len(files), flush=True)

    groups = {}
    for line in lines:
        key = folded(line["concept"])
        item = groups.setdefault(key, {"concept": line["concept"], "family": line["family"], "group": line["group"], "basis": line["basis"], "unit_codes": set(), "invoice_ids": set(), "line_count": 0, "amount": 0.0, "variants": Counter()})
        item["unit_codes"].add(line["unit_code"])
        item["invoice_ids"].add(line["invoice"])
        item["line_count"] += 1
        item["amount"] += line["amount"]
        item["variants"][line["description"]] += 1
    concepts = []
    for item in groups.values():
        concepts.append({
            "concept": item["concept"], "family": item["family"], "group": item["group"], "basis": item["basis"],
            "unit_codes": sorted(x for x in item["unit_codes"] if x),
            "invoices": len(item["invoice_ids"]), "lines": item["line_count"],
            "amount": round(item["amount"], 2), "variants": item["variants"].most_common(3),
        })
    concepts.sort(key=lambda x: (-x["lines"], x["concept"].casefold()))
    payload = {"source_files": len(files), "lines": lines, "concepts": concepts, "issues": issues, "invoices_without_lines": invoices_without_lines, "header_patterns": [{"row": a, "description_col": b, "amount_col": c, "count": n} for (a, b, c), n in header_patterns.most_common()]}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"files": len(files), "lines": len(lines), "concepts": len(concepts), "issues": len(issues), "without_lines": len(invoices_without_lines), "families": dict(Counter(x["family"] for x in lines)), "headers": payload["header_patterns"][:10]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
