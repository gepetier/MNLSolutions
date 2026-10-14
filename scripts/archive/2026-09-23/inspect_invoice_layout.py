import os
from pathlib import Path
from openpyxl import load_workbook

root = Path(os.environ["MNL_INVOICE_SOURCE_DIR"])
files = sorted(p for p in root.rglob("*.xlsx") if "factura" in p.name.lower() or "invoice" in p.name.lower())
for index in [650, len(files)-1]:
    path = files[index]
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    print("\nFILE", index, "sheets", len(book.worksheets), "shape", sheet.max_row, sheet.max_column)
    for rownum, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 75), values_only=False), 1):
        cells = []
        for cell in row[:18]:
            if cell.value is None:
                continue
            val = str(cell.value).replace("\n", " ").strip()
            if len(val) > 75:
                val = val[:72] + "..."
            cells.append(f"{cell.column_letter}:{val}")
        if cells:
            print(rownum, " | ".join(cells))
