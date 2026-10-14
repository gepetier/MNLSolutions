from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from itertools import islice
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SENSITIVE_DATA_DIR = Path(os.environ.get("MNL_SENSITIVE_DATA_DIR", r"Z:\Adri\Dades MNLSolutions"))
PYDEPS = ROOT / "pydeps"
if PYDEPS.exists():
    sys.path.insert(0, str(PYDEPS))

QUOTE_PREFIXES = (
    "presp",
    "comanda",
    "pedido",
)
SKIP_PREFIXES = (
    "no valid",
    "no_val",
    "nou format",
)

DOCUMENT_TYPES = ("quotes", "invoices")


def compact_spaces(value: str) -> str:
    text = str(value or "").strip()
    if "\u00c3" in text or "\u00c2" in text:
        try:
            text = text.encode("cp1252").decode("utf-8")
        except UnicodeError:
            pass
    if any(marker in text for marker in ("Ã", "Â", "â")):
        try:
            text = text.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    text = text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text)


def normalize_name(value: str) -> str:
    return compact_spaces(value).strip(" ,;:-_")


def normalize_name_key(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize_name(value).upper())


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D+", "", value or "")
    if digits.startswith("34") and len(digits) == 11:
        digits = digits[2:]
    return digits


def format_phone(value: str) -> str:
    digits = normalize_phone(value)
    if len(digits) == 9:
        return f"+34 {digits[:3]} {digits[3:6]} {digits[6:]}"
    if digits.startswith("00") and len(digits) > 10:
        return f"+{digits[2:]}"
    if digits.startswith("34") and len(digits) == 11:
        return f"+34 {digits[2:5]} {digits[5:8]} {digits[8:]}"
    return digits


def normalize_vat(value: str) -> str:
    value = compact_spaces(value).upper()
    value = re.sub(r"^(N\.?I\.?F\.?|CIF|VAT|NIF)\s*:?\s*", "", value)
    value = re.sub(r"[^A-Z0-9]", "", value.replace("ES", "", 1))
    if not re.search(r"\d", value):
        return ""
    if len(value) < 8 or len(value) > 10:
        return ""
    if not re.match(r"^[A-Z0-9][A-Z0-9]{7,9}$", value):
        return ""
    return value


def extract_vat(value: str) -> str:
    for match in re.finditer(r"(?:N\.?I\.?F\.?|CIF|VAT|NIF)\s*:?\s*([A-Z0-9 .-]{8,14})", value or "", re.I):
        vat = normalize_vat(match.group(1))
        if vat:
            return vat
    return normalize_vat(value)


def extract_emails(value: str) -> list[str]:
    emails = re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", value or "")
    unique = []
    for email in emails:
        email = email.lower()
        if email not in unique:
            unique.append(email)
    return unique


def extract_email(value: str) -> str:
    emails = extract_emails(value)
    return emails[0] if emails else ""


def extract_phones(value: str) -> list[str]:
    text = re.sub(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", " ", value or "")
    phones = []
    label_matches = re.finditer(
        r"(?:Tlf|Tel|Tel\.|Telefono|Tel.fono|Movil|M.vil|Mobile)\s*:?\s*([+()\d .-]{6,})",
        text,
        re.I,
    )
    for match in label_matches:
        digits = normalize_phone(match.group(1))
        if len(digits) >= 9 and digits not in phones:
            phones.append(digits)
    for match in re.finditer(r"(?<!\d)(?:\+?34[\s.-]?)?([6789]\d(?:[\s.-]?\d){7})(?!\d)", text):
        digits = normalize_phone(match.group(0))
        if len(digits) == 9 and digits not in phones:
            phones.append(digits)
    return phones


def extract_phone(value: str) -> str:
    phones = extract_phones(value)
    return phones[0] if phones else ""


def normalize_street(value: str) -> str:
    text = compact_spaces(value)
    text = re.sub(r"^(Direccion|Direcci.n|Adre.a|Domicilio)\s*:?\s*", "", text, flags=re.I)
    text = text.replace("C\\", "C/")
    text = re.sub(r"\s*,\s*", ", ", text)
    text = re.sub(r"\s*/\s*", "/", text)
    text = text.strip(" ,;")
    if re.fullmatch(r"C/?", text, flags=re.I):
        return ""
    return text


PROVINCE_ALIASES = {
    "BCN": "Barcelona",
    "BARNA": "Barcelona",
    "BARCELONA": "Barcelona",
    "GIRONA": "Girona",
    "GERONA": "Girona",
    "LLEIDA": "Lleida",
    "LERIDA": "Lleida",
    "TARRAGONA": "Tarragona",
}


def normalize_province(value: str) -> str:
    text = compact_spaces(value).strip(" .")
    return PROVINCE_ALIASES.get(text.upper(), text)


def parse_city_line(value: str) -> dict:
    text = compact_spaces(value).strip(" ,;")
    text = re.sub(r"\s*-\s*", " - ", text)
    if not text:
        return {"cityLine": "", "zip": "", "city": "", "province": "", "country": ""}

    zip_code = ""
    match = re.match(r"^(\d{5})\s*(?:-|,)?\s*(.*)$", text)
    if match:
        zip_code = match.group(1)
        text = match.group(2).strip()

    province = ""
    match = re.search(r"\(([^)]+)\)\s*$", text)
    if match:
        province = normalize_province(match.group(1))
        text = text[: match.start()].strip(" ,;-")

    city = re.sub(r"\s+", " ", text).strip(" ,;-")
    city_line = " - ".join(part for part in [zip_code, city] if part)
    if province:
        city_line = f"{city_line} ({province})" if city_line else province
    country = "Spain" if zip_code or province else ""
    return {
        "cityLine": city_line,
        "zip": zip_code,
        "city": city,
        "province": province,
        "country": country,
    }


def build_contact_quality(contact: dict) -> dict:
    quality = {
        "hasEmail": bool(contact.get("email")),
        "hasPhone": bool(contact.get("phoneDigits") or contact.get("phone")),
        "hasVat": bool(contact.get("vat")),
        "hasStreet": bool(contact.get("street")),
        "hasCity": bool(contact.get("city")),
        "hasZip": bool(contact.get("zip")),
    }
    quality["hasAddress"] = quality["hasStreet"] or quality["hasCity"]
    quality["addressComplete"] = quality["hasStreet"] and quality["hasCity"]
    quality["contactable"] = quality["hasEmail"] or quality["hasPhone"]
    quality["missing"] = [
        label
        for label, present in [
            ("email", quality["hasEmail"]),
            ("phone", quality["hasPhone"]),
            ("vat", quality["hasVat"]),
            ("street", quality["hasStreet"]),
            ("city", quality["hasCity"]),
        ]
        if not present
    ]
    quality["score"] = sum(
        1
        for key in ["hasEmail", "hasPhone", "hasVat", "hasStreet", "hasCity", "hasZip"]
        if quality[key]
    )
    return quality


def document_ref_from_name(path: Path) -> str:
    name = compact_spaces(path.name)
    match = re.search(r"MNL[-._ ]?(\d{2})[\u00b7.\-_\s]*(\d{4})[-._ ]?([A-Z])?", name, re.I)
    if not match:
        return ""
    suffix = f"-{match.group(3).upper()}" if match.group(3) else ""
    return f"MNL-{match.group(1)}\u00b7{match.group(2)}{suffix}"


def infer_name_from_path(path: Path) -> str:
    for parent in [path.parent, *path.parents]:
        folder_name = compact_spaces(parent.name)
        if not re.match(r"^MNL[-_ ]?\d{2}[\u00b7._-]?\d{4}", folder_name, re.I):
            continue
        inferred = re.sub(r"^MNL[-_ ]?\d{2}[\u00b7._-]?\d{4}[-_ ]?[A-Z]?[_ -]*", "", folder_name, flags=re.I)
        inferred = re.sub(r"\([^)]*\)", "", inferred).strip(" _-")
        return normalize_name(inferred)
    return ""


def date_to_iso(value) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = compact_spaces(value)
    if re.fullmatch(r"\d{5}(?:\.0+)?", text):
        try:
            return (date(1899, 12, 30) + timedelta(days=int(float(text)))).isoformat()
        except ValueError:
            pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def read_workbook(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=False)
        sheet = book.sheet_by_index(0)

        def cell(row: int, col: int) -> str:
            if row >= sheet.nrows or col >= sheet.ncols:
                return ""
            value = sheet.cell_value(row, col)
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            return compact_spaces(value)

        return cell

    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    sheet = book.worksheets[0]

    def cell(row: int, col: int) -> str:
        value = sheet.cell(row=row + 1, column=col + 1).value
        return compact_spaces(value)

    return cell


def candidate_files(root: Path, document_type: str):
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}:
            continue
        name = path.name.lower().strip()
        if name.startswith("~$") or name.startswith(SKIP_PREFIXES):
            continue
        if document_type == "quotes" and not name.startswith(QUOTE_PREFIXES):
            continue
        if document_type == "invoices" and not any(token in name for token in ("factura", "invoice")):
            continue
        yield path


def find_labeled_value(cell, labels: tuple[str, ...], max_rows: int = 20, max_cols: int = 16) -> tuple[str, int, int]:
    normalized_labels = tuple(normalize_name_key(label) for label in labels)
    for row in range(max_rows):
        for col in range(max_cols):
            value = compact_spaces(cell(row, col))
            value_key = normalize_name_key(value)
            if not value_key or not any(label in value_key for label in normalized_labels):
                continue
            for next_col in range(col + 1, max_cols):
                candidate = compact_spaces(cell(row, next_col))
                if candidate:
                    return candidate, row, col
            for next_row in range(row + 1, min(row + 7, max_rows)):
                candidate = compact_spaces(cell(next_row, col))
                if candidate:
                    return candidate, row, col
            return "", row, col
    return "", -1, -1


def invoice_ref_from_name(path: Path) -> str:
    match = re.search(r"(?:^|[_ -])([A-Z]\d{2}-\d+)(?:[_ -]|$)", path.stem, re.I)
    return match.group(1).upper() if match else ""


def infer_invoice_name_from_path(path: Path) -> str:
    stem = compact_spaces(path.stem)
    match = re.search(
        r"(?:factura|invoice)(?:\s+proforma)?[_ -]*(?:[A-Z]\d{2}-\d+)?[_ -]+(.+)$",
        stem,
        re.I,
    )
    if not match:
        return ""
    return normalize_name(match.group(1))


def extract_invoice_contact(path: Path, cell) -> dict | None:
    name, billing_row, billing_col = find_labeled_value(cell, ("Facturar a", "Bill to", "Cliente"))
    if billing_row < 0:
        return None

    customer_lines = []
    for row in range(billing_row + 1, billing_row + 7):
        parts = [cell(row, col) for col in range(billing_col, min(billing_col + 5, 16))]
        text = compact_spaces(" ".join(str(part) for part in parts if part))
        if text:
            customer_lines.append(text)
    customer_text = " ".join(customer_lines)

    name = normalize_name(name)
    street = normalize_street(customer_lines[1] if len(customer_lines) > 1 else "")
    city_parts = parse_city_line(customer_lines[2] if len(customer_lines) > 2 else "")
    vat = extract_vat(customer_text)
    emails = extract_emails(customer_text)
    phone_digits_list = extract_phones(customer_text)
    phone_digits = phone_digits_list[0] if phone_digits_list else ""

    bad_name_markers = ("MNL ", "HUURRE ", "FACTURA", "INVOICE", "FACTURAR A", "MATERIAL")
    address_prefixes = ("C/", "C\\", "AVDA", "AV.", "CARRETERA", "CTRA", "POL.", "POLIGONO")
    looks_like_address = name.upper().startswith(address_prefixes)
    looks_like_internal_code = bool(re.fullmatch(r"[A-Z]{1,3}\d{3,}", name.upper()))
    if not name or name.upper().startswith(bad_name_markers) or looks_like_address or looks_like_internal_code:
        inferred = infer_invoice_name_from_path(path)
        if inferred:
            name = inferred
    if not name or name.upper().startswith(bad_name_markers):
        return None

    invoice_ref, _ref_row, _ref_col = find_labeled_value(cell, ("Nº de Factura", "Nº Factura", "No Factura", "Numero Factura"))
    invoice_ref = compact_spaces(invoice_ref) or invoice_ref_from_name(path)
    invoice_date, _date_row, _date_col = find_labeled_value(cell, ("Fecha", "Data", "Date"))
    invoice_date = date_to_iso(invoice_date)
    reference, _reference_row, _reference_col = find_labeled_value(cell, ("Referencia", "Reference"))

    contact = {
        "name": name,
        "vat": vat,
        "email": emails[0] if emails else "",
        "emails": emails,
        "phone": format_phone(phone_digits) if phone_digits else "",
        "phoneDigits": phone_digits,
        "phones": [format_phone(phone) for phone in phone_digits_list],
        "street": street,
        "street2": "",
        **city_parts,
        "quoteRefs": [],
        "invoiceRefs": [invoice_ref] if invoice_ref else [],
        "documentRefs": [invoice_ref] if invoice_ref else [],
        "firstQuoteDate": "",
        "lastQuoteDate": "",
        "firstDocumentDate": invoice_date,
        "lastDocumentDate": invoice_date,
        "references": [compact_spaces(reference)] if reference else [],
        "sourceFiles": [str(path)],
        "sourceTypes": ["invoice"],
    }
    contact["quality"] = build_contact_quality(contact)
    return contact


def extract_contact(path: Path, document_type: str) -> dict | None:
    try:
        cell = read_workbook(path)
    except Exception as error:
        return {
            "sourceFile": str(path),
            "error": str(error),
        }

    if document_type == "invoices":
        return extract_invoice_contact(path, cell)

    # MNL/Huurre issuer blocks sit at columns A-D. The customer block starts at H.
    def line(row: int, start: int = 7, end: int = 12) -> str:
        parts = [cell(row, col) for col in range(start, end)]
        return compact_spaces(" ".join(part for part in parts if part))

    customer_lines = [line(row, 7, 12) for row in range(6, 13)]
    customer_text = " ".join(part for part in customer_lines if part)

    name = normalize_name(line(6, 7, 10))
    street = normalize_street(line(7, 7, 10))
    city_parts = parse_city_line(line(8, 7, 10))
    vat = extract_vat(line(9, 7, 12)) or extract_vat(customer_text)
    emails = extract_emails(customer_text)
    phone_digits_list = extract_phones(customer_text)
    phone_digits = phone_digits_list[0] if phone_digits_list else ""
    address_prefixes = ("C/", "C\\", "AVDA", "AV.", "CARRETERA", "CTRA", "POL.", "POLIGONO", "POLIGONO")
    address_name = name.upper().startswith(address_prefixes)
    if address_name:
        inferred = infer_name_from_path(path)
        if inferred:
            name = inferred

    bad_name_markers = ("MNL ", "HUURRE ", "P R E S", "PRESSUPOST", "PRESUPUESTO", "3.-PRESSUPOST")
    if not name or name.upper().startswith(bad_name_markers) or "PRESSUPOST MNL" in name.upper():
        return None

    ref = document_ref_from_name(path) or compact_spaces(cell(3, 11))
    quote_date = date_to_iso(cell(4, 11))
    reference = compact_spaces(cell(2, 11))

    if not any([street, city_parts["cityLine"], vat, emails, phone_digits]):
        inferred = infer_name_from_path(path)
        if inferred and len(inferred) > len(name):
            name = normalize_name(inferred)

    if not name or name.upper().startswith(bad_name_markers) or "PRESSUPOST MNL" in name.upper():
        return None

    contact = {
        "name": name,
        "vat": vat,
        "email": emails[0] if emails else "",
        "emails": emails,
        "phone": format_phone(phone_digits) if phone_digits else "",
        "phoneDigits": phone_digits,
        "phones": [format_phone(phone) for phone in phone_digits_list],
        "street": street,
        "street2": "",
        **city_parts,
        "quoteRefs": [ref] if ref else [],
        "invoiceRefs": [],
        "documentRefs": [ref] if ref else [],
        "firstQuoteDate": quote_date,
        "lastQuoteDate": quote_date,
        "firstDocumentDate": quote_date,
        "lastDocumentDate": quote_date,
        "references": [reference] if reference else [],
        "sourceFiles": [str(path)],
        "sourceTypes": ["quote"],
    }
    contact["quality"] = build_contact_quality(contact)
    return contact


def dedupe_key(contact: dict) -> str:
    if contact.get("vat"):
        return f"vat:{contact['vat']}"
    if contact.get("email"):
        return f"email:{contact['email']}"
    if contact.get("phoneDigits"):
        return f"phone:{contact['phoneDigits']}"
    return "name:" + normalize_name_key(contact.get("name", ""))


def merge_unique(target: dict, source: dict, key: str):
    values = target.setdefault(key, [])
    for value in source.get(key, []):
        if value and value not in values:
            values.append(value)


def refresh_primary_fields(contact: dict):
    contact["email"] = (contact.get("emails") or [contact.get("email", "")])[0] if (contact.get("emails") or contact.get("email")) else ""
    contact["phone"] = (contact.get("phones") or [contact.get("phone", "")])[0] if (contact.get("phones") or contact.get("phone")) else ""
    contact["phoneDigits"] = normalize_phone(contact.get("phoneDigits") or contact.get("phone"))
    contact["quality"] = build_contact_quality(contact)


def merge_contact(target: dict, source: dict):
    for key in ["vat", "street", "street2", "cityLine", "zip", "city", "province", "country"]:
        if not target.get(key) and source.get(key):
            target[key] = source[key]
    for key in ["emails", "phones", "quoteRefs", "invoiceRefs", "documentRefs", "references", "sourceFiles", "sourceTypes"]:
        merge_unique(target, source, key)
    if not target.get("phoneDigits") and source.get("phoneDigits"):
        target["phoneDigits"] = source["phoneDigits"]
    dates = [
        d
        for d in [
            target.get("firstQuoteDate"),
            target.get("lastQuoteDate"),
            source.get("firstQuoteDate"),
            source.get("lastQuoteDate"),
            target.get("firstDocumentDate"),
            target.get("lastDocumentDate"),
            source.get("firstDocumentDate"),
            source.get("lastDocumentDate"),
        ]
        if d
    ]
    iso_dates = [d for d in dates if re.match(r"^\d{4}-\d{2}-\d{2}$", d)]
    if iso_dates:
        target["firstQuoteDate"] = min(iso_dates)
        target["lastQuoteDate"] = max(iso_dates)
        target["firstDocumentDate"] = min(iso_dates)
        target["lastDocumentDate"] = max(iso_dates)
    refresh_primary_fields(target)


def collapse_by_name(contacts_by_key: dict) -> dict:
    by_name = defaultdict(list)
    for key, contact in contacts_by_key.items():
        by_name[normalize_name_key(contact.get("name", ""))].append((key, contact))

    collapsed = {}
    consumed = set()
    for _name_key, items in by_name.items():
        rich_items = [
            (key, contact)
            for key, contact in items
            if contact.get("vat") or contact.get("email") or contact.get("phone")
        ]
        if not rich_items or len(items) == 1:
            for key, contact in items:
                if key not in consumed:
                    refresh_primary_fields(contact)
                    collapsed[key] = contact
            continue

        target_key, target = rich_items[0]
        for key, contact in items:
            if key == target_key:
                continue
            merge_contact(target, contact)
            consumed.add(key)
        refresh_primary_fields(target)
        collapsed[target_key] = target
        consumed.add(target_key)

    for key, contact in contacts_by_key.items():
        if key not in consumed and key not in collapsed:
            refresh_primary_fields(contact)
            collapsed[key] = contact
    return collapsed


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("--document-type", choices=DOCUMENT_TYPES, default="quotes")
    parser.add_argument("--output", type=Path, default=SENSITIVE_DATA_DIR / "contacts.json")
    parser.add_argument("--limit", type=int, help="Processa només els primers fitxers candidats (per validació).")
    args = parser.parse_args()

    contacts_by_key = {}
    errors = []
    scanned = 0
    for path in islice(candidate_files(args.root, args.document_type), args.limit):
        scanned += 1
        contact = extract_contact(path, args.document_type)
        if not contact:
            continue
        if contact.get("error"):
            errors.append(contact)
            continue
        key = dedupe_key(contact)
        if key in contacts_by_key:
            merge_contact(contacts_by_key[key], contact)
        else:
            contact["id"] = re.sub(r"[^a-z0-9]+", "-", key.lower()).strip("-")
            contact["odooId"] = None
            contact["selected"] = True
            refresh_primary_fields(contact)
            contacts_by_key[key] = contact

    contacts_by_key = collapse_by_name(contacts_by_key)
    contacts = sorted(contacts_by_key.values(), key=lambda item: item.get("name", "").upper())
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceRoot": str(args.root),
        "documentType": args.document_type,
        "scannedFiles": scanned,
        "contactCount": len(contacts),
        "errorCount": len(errors),
        "contacts": contacts,
        "errors": errors,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {key: payload[key] for key in ["generatedAt", "scannedFiles", "contactCount", "errorCount"]},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
