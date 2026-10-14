from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "pydeps"))
sys.path.insert(0, str(ROOT / "scripts"))

from extract_contacts_from_quotes import (  # noqa: E402
    candidate_files,
    compact_spaces,
    extract_vat,
    infer_invoice_name_from_path,
    invoice_ref_from_name,
    normalize_name,
    date_to_iso,
)


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def as_number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip().replace("€", "").replace("EUR", "").replace("\xa0", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def iso_date(value: object) -> str:
    text = date_to_iso(value)
    return text if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text) else ""


def find_labeled_value(rows: list[tuple], labels: set[str]) -> object | None:
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            if not isinstance(value, str) or not any(label in normalized(value) for label in labels):
                continue
            for candidate in row[column + 1 : column + 6]:
                if candidate not in (None, ""):
                    return candidate
            for next_row in rows[row_index + 1 : row_index + 7]:
                if column < len(next_row) and next_row[column] not in (None, ""):
                    return next_row[column]
    return None


def customer_details(rows: list[tuple]) -> tuple[str, str]:
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            if not isinstance(value, str) or not any(
                label in normalized(value) for label in {"facturar a", "bill to"}
            ):
                continue
            name = None
            for candidate in row[column + 1 : column + 6]:
                if candidate not in (None, ""):
                    name = candidate
                    break
            if name is None:
                for next_row in rows[row_index + 1 : row_index + 2]:
                    if column < len(next_row) and next_row[column] not in (None, ""):
                        name = next_row[column]
                        break
            customer_lines = []
            for customer_row in rows[row_index + 1 : row_index + 7]:
                parts = customer_row[column : column + 5]
                line = compact_spaces(" ".join(str(part) for part in parts if part not in (None, "")))
                if line:
                    customer_lines.append(line)
            return normalize_name(name), extract_vat(" ".join(customer_lines))
    return "", ""


def invoice_total(rows: list[tuple]) -> float | None:
    candidates: list[tuple[int, float]] = []
    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            if not isinstance(value, str):
                continue
            label = normalized(value)
            if not label.startswith("total") or "subtotal" in label or label in {
                "totales",
                "total unidades",
                "total unidades ml",
            }:
                continue
            numbers = [as_number(cell) for cell in row[column + 1 : column + 6]]
            numbers = [number for number in numbers if number is not None]
            if numbers:
                candidates.append((row_number, numbers[-1]))
    final_rows = [candidate for candidate in candidates if candidate[0] > 18]
    return (final_rows[-1] if final_rows else candidates[-1])[1] if candidates else None


def currency_code(rows: list[tuple]) -> str:
    text = " ".join(str(cell) for row in rows for cell in row if isinstance(cell, str)).upper()
    if "EUR" in text or "€" in text:
        return "EUR"
    match = re.search(r"\b(USD|GBP|CHF|CAD|AUD)\b", text)
    return match.group(1) if match else "EUR"


def file_key(path: Path, root: Path) -> str:
    relative = path.relative_to(root).as_posix().casefold()
    return hashlib.sha256(relative.encode("utf-8")).hexdigest()


def prepare(root: Path) -> list[dict]:
    root = root.resolve()
    pdf_by_parent_and_stem: dict[tuple[Path, str], list[Path]] = {}
    for pdf in root.rglob("*.pdf"):
        key = (pdf.parent, normalized(pdf.stem))
        pdf_by_parent_and_stem.setdefault(key, []).append(pdf)

    documents = []
    for path in candidate_files(root, "invoices"):
        filename_key = path.name.casefold()
        if "proforma" in filename_key or "plantilla" in filename_key:
            continue
        year_match = re.search(r"20\d{2}", path.parent.name)
        if not year_match:
            continue
        document: dict = {
            "source_key": file_key(path, root),
            "source_filename": path.name,
            "source_year": int(year_match.group()),
            "invoice_number": invoice_ref_from_name(path),
            "invoice_date": "",
            "due_date": "",
            "customer_name": "",
            "customer_vat": "",
            "partner_reference": "",
            "amount_total": None,
            "amount_valid": False,
            "currency_code": "EUR",
            "pdf_filename": "",
            "review_status": "review",
            "extraction_note": "",
        }
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.worksheets[0]
            rows = list(
                sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 300),
                    max_col=min(sheet.max_column, 20),
                    values_only=True,
                )
            )
            document["customer_name"], document["customer_vat"] = customer_details(rows)
            if not document["invoice_number"]:
                document["invoice_number"] = compact_spaces(
                    find_labeled_value(rows, {"n de factura", "n factura", "numero factura", "no factura"})
                )
            if not document["invoice_date"]:
                document["invoice_date"] = iso_date(
                    find_labeled_value(rows, {"fecha", "fecha factura", "data", "date"})
                )
            document["due_date"] = iso_date(
                find_labeled_value(rows, {"vence", "venciment", "vencimento", "due date"})
            )
            if not document["customer_name"]:
                document["customer_name"] = compact_spaces(infer_invoice_name_from_path(path))
            document["partner_reference"] = compact_spaces(
                find_labeled_value(rows, {"referencia", "reference"})
            )
            document["amount_total"] = invoice_total(rows)
            document["amount_valid"] = document["amount_total"] is not None
            document["currency_code"] = currency_code(rows)
            workbook.close()
        except Exception as error:
            document["extraction_note"] = "Excel no llegible: " + compact_spaces(error)

        pdfs = pdf_by_parent_and_stem.get((path.parent, normalized(path.stem)), [])
        if len(pdfs) == 1:
            document["pdf_filename"] = pdfs[0].name
        elif len(pdfs) > 1:
            document["extraction_note"] = (document["extraction_note"] + " PDF ambigu.").strip()

        has_core_fields = bool(document["invoice_number"] and document["invoice_date"] and document["customer_name"])
        document["review_status"] = "ready" if has_core_fields and document["amount_valid"] else "review"
        if not document["invoice_number"]:
            document["extraction_note"] = (document["extraction_note"] + " Falta número de factura.").strip()
        if not document["invoice_date"]:
            document["extraction_note"] = (document["extraction_note"] + " Falta data de factura.").strip()
        if not document["customer_name"]:
            document["extraction_note"] = (document["extraction_note"] + " Falta client.").strip()
        if not document["amount_valid"]:
            document["extraction_note"] = (document["extraction_note"] + " Total no fiable.").strip()
        documents.append(document)
    return documents


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: prepare_historical_invoice_archive.py <invoice-folder>")
    print(json.dumps(prepare(Path(sys.argv[1])), ensure_ascii=False))
